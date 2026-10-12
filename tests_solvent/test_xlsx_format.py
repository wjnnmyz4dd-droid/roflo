"""The ``.xlsx`` format layer, tested on its own and against another library.

This module has its own suite rather than being covered incidentally through
the capability that uses it, for one reason: the worker and the verifier both
read spreadsheets through it, exactly as both read CSV through the standard
library's ``csv``. So a bug *here* fools both sides of a two-sided design,
which is the one way that design can fail quietly.

Two kinds of test:

**Refusals.** A file this module cannot read correctly must raise rather than
return a best effort. An approximate read of a client's spreadsheet is a
confident wrong answer, and the client cannot tell by looking.

**Cross-validation against openpyxl.** Our own round-trip proves we agree with
ourselves, which is worth very little. These tests ask an independent
implementation to read what we wrote and write what we read. ``openpyxl`` is a
**test-only** tool: Solvent imports only the standard library, and the Windows
installer builds a virtual environment with no packages, so nothing shipped
depends on it. The tests skip when it is absent, and that is a real reduction
in evidence rather than a formality -- which is why the skip is reported
rather than silent.

The LibreOffice binary in this container cannot load even a plain CSV, so it
was not usable as a second validator. Recorded here so nobody repeats the
attempt expecting it to work.
"""

from __future__ import annotations

import datetime as dt
import io
import unittest
import zipfile

from solvent import xlsx

try:
    import openpyxl
except ImportError:                                  # pragma: no cover
    openpyxl = None

NEEDS_OPENPYXL = unittest.skipIf(
    openpyxl is None,
    "openpyxl is not installed. It is a test-only cross-validator, never a "
    "Solvent dependency -- but without it the only evidence that our files "
    "are readable by anything else is gone, so this is a weaker run")


def package(parts: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as out:
        for name, text in parts.items():
            out.writestr(name, text)
    return buffer.getvalue()


MINIMAL = {
    "[Content_Types].xml": "<Types/>",
    "xl/workbook.xml":
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
        '2006/main"><sheets><sheet name="S" sheetId="1"/></sheets></workbook>',
    "xl/worksheets/sheet1.xml":
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
        '2006/main"><sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>hi'
        "</t></is></c></row></sheetData></worksheet>",
}


class AFileItCannotReadIsRefused(unittest.TestCase):

    def test_something_that_is_not_a_zip(self):
        with self.assertRaises(xlsx.UnsupportedWorkbook) as caught:
            xlsx.read_workbook(b"this is not a spreadsheet")
        self.assertIn("not a readable", str(caught.exception))

    def test_a_zip_with_no_workbook_part(self):
        with self.assertRaises(xlsx.UnsupportedWorkbook) as caught:
            xlsx.read_workbook(package({"hello.txt": "hi"}))
        self.assertIn("workbook.xml", str(caught.exception))

    def test_a_macro_enabled_workbook(self):
        """The macros would not run, and a file whose behaviour depends on
        code we ignore cannot be processed correctly."""
        with self.assertRaises(xlsx.UnsupportedWorkbook) as caught:
            xlsx.read_workbook(package(
                {**MINIMAL, "xl/vbaProject.bin": "MZ"}))
        self.assertIn("macro", str(caught.exception).lower())

    def test_a_workbook_linking_to_other_files(self):
        """Its values depend on files we do not have, so a total computed
        from it could be wrong without anything looking wrong."""
        with self.assertRaises(xlsx.UnsupportedWorkbook) as caught:
            xlsx.read_workbook(package(
                {**MINIMAL, "xl/externalLinks/externalLink1.xml": "<x/>"}))
        self.assertIn("other files", str(caught.exception))

    def test_a_worksheet_whose_xml_is_broken(self):
        with self.assertRaises(xlsx.UnsupportedWorkbook):
            xlsx.read_workbook(package(
                {**MINIMAL, "xl/worksheets/sheet1.xml": "<worksheet>"}))

    def test_a_sheet_declaring_more_cells_than_we_will_process(self):
        """A crafted file can declare an enormous sheet, and a worker that
        tried to honour it would exhaust the machine. Raising rather than
        truncating: half a client's data silently processed is worse."""
        cells = "".join(f'<c r="A{i}"><v>{i}</v></c>' for i in range(1, 60))
        rows = "".join(f'<row r="{i}">{cells}</row>' for i in range(1, 60))
        with self.assertRaises(xlsx.UnsupportedWorkbook) as caught:
            xlsx.read_workbook(package({
                **MINIMAL,
                "xl/worksheets/sheet1.xml":
                    '<worksheet xmlns="http://schemas.openxmlformats.org/'
                    f'spreadsheetml/2006/main"><sheetData>{rows}</sheetData>'
                    "</worksheet>"}), max_cells=100)
        self.assertIn("beyond what", str(caught.exception))

    def test_an_xml_entity_attack_does_not_read_the_filesystem(self):
        """A client's spreadsheet that tries to read /etc/passwd through a
        DOCTYPE is an attack, not a file to be helpful about."""
        hostile = (
            '<?xml version="1.0"?>'
            '<!DOCTYPE t [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/'
            'spreadsheetml/2006/main"><sheetData><row r="1">'
            '<c r="A1" t="inlineStr"><is><t>&x;</t></is></c>'
            "</row></sheetData></worksheet>")
        try:
            book = xlsx.read_workbook(package(
                {**MINIMAL, "xl/worksheets/sheet1.xml": hostile}))
        except xlsx.UnsupportedWorkbook:
            return          # refused outright, which is the best outcome
        for sheet in book.sheets:
            for row in sheet.rows:
                for cell in row:
                    self.assertNotIn("root:", str(cell.value or ""))

    def test_asking_for_a_sheet_that_is_not_there(self):
        book = xlsx.read_workbook(package(MINIMAL))
        with self.assertRaises(xlsx.UnsupportedWorkbook) as caught:
            book.sheet("Nope")
        self.assertIn("'S'", str(caught.exception))


class ValuesAreReadAsWhatTheyAre(unittest.TestCase):

    def test_a_shared_string_index_outside_the_table_is_empty_not_invented(self):
        """Empty is the honest reading of a broken file. Inventing text would
        put words in a client's document that nobody wrote."""
        book = xlsx.read_workbook(package({
            **MINIMAL,
            "xl/sharedStrings.xml":
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
                '2006/main"><si><t>only one</t></si></sst>',
            "xl/worksheets/sheet1.xml":
                '<worksheet xmlns="http://schemas.openxmlformats.org/'
                'spreadsheetml/2006/main"><sheetData><row r="1">'
                '<c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>9</v></c>'
                "</row></sheetData></worksheet>"}))
        self.assertEqual(book.first.values()[0], ["only one", ""])

    def test_an_excel_error_value_stays_text_and_is_not_coerced(self):
        book = xlsx.read_workbook(package({
            **MINIMAL,
            "xl/worksheets/sheet1.xml":
                '<worksheet xmlns="http://schemas.openxmlformats.org/'
                'spreadsheetml/2006/main"><sheetData><row r="1">'
                '<c r="A1" t="e"><f>1/0</f><v>#DIV/0!</v></c>'
                "</row></sheetData></worksheet>"}))
        cell = book.first.rows[0][0]
        self.assertEqual(cell.value, "#DIV/0!")
        self.assertEqual(cell.formula, "1/0")

    def test_the_sheet_part_is_resolved_through_the_relationships(self):
        """Not by guessing sheet1.xml, sheet2.xml and so on. The mapping is
        not positional, and a workbook that has had sheets deleted and
        re-added does not follow that pattern -- guessing reads the wrong
        sheet, which is the worst kind of wrong because every value in it is
        valid."""
        rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        book = xlsx.read_workbook(package({
            "[Content_Types].xml": "<Types/>",
            "xl/workbook.xml":
                '<workbook xmlns="http://schemas.openxmlformats.org/'
                f'spreadsheetml/2006/main" xmlns:r="{rel}"><sheets>'
                '<sheet name="Wanted" sheetId="7" r:id="rId9"/>'
                "</sheets></workbook>",
            "xl/_rels/workbook.xml.rels":
                '<Relationships xmlns="http://schemas.openxmlformats.org/'
                'package/2006/relationships"><Relationship Id="rId9" '
                f'Type="{rel}/worksheet" Target="worksheets/odd_name.xml"/>'
                "</Relationships>",
            "xl/worksheets/odd_name.xml":
                '<worksheet xmlns="http://schemas.openxmlformats.org/'
                'spreadsheetml/2006/main"><sheetData><row r="1">'
                '<c r="A1"><v>42</v></c></row></sheetData></worksheet>',
        }))
        self.assertEqual([s.name for s in book.sheets], ["Wanted"])
        self.assertEqual(book.first.values(), [[42]])

    def test_excels_phantom_leap_day_is_not_turned_into_a_date(self):
        """Serial 60 is the 29th of February 1900, which did not exist.
        Returned as the number it is rather than a date nobody can name."""
        self.assertEqual(xlsx._serial_to_date(60), 60)

    def test_dates_either_side_of_the_phantom_day_are_right(self):
        """Getting this wrong moves every date in a client's file by one,
        which is easy to do and hard for them to spot."""
        for serial, expected in ((1, "1900-01-01"), (59, "1900-02-28"),
                                 (61, "1900-03-01"), (44927, "2023-01-01"),
                                 (45291, "2023-12-31")):
            with self.subTest(serial=serial):
                self.assertEqual(str(xlsx._serial_to_date(serial)), expected)

    def test_a_date_survives_a_round_trip_through_a_serial(self):
        for day in (dt.date(1900, 1, 1), dt.date(1900, 2, 28),
                    dt.date(1900, 3, 1), dt.date(2026, 10, 12),
                    dt.date(2099, 12, 31)):
            with self.subTest(day=day):
                serial = xlsx._date_to_serial(day)
                self.assertEqual(xlsx._serial_to_date(serial), day)

    def test_a_number_format_mentioning_a_month_in_quotes_is_not_a_date(self):
        """``"May"0.00`` displays the word May beside a number. A format
        matcher that looked at quoted text would turn every such number into
        a date."""
        self.assertFalse(xlsx._is_date_format('"May"0.00'))
        self.assertFalse(xlsx._is_date_format('[Red]0.00'))
        self.assertTrue(xlsx._is_date_format("yyyy-mm-dd"))


class WritingProducesAFileAndTheSameFileTwice(unittest.TestCase):

    def test_the_same_input_produces_the_same_bytes(self):
        """A deliverable's digest is what the audit record names. A file whose
        digest changed on every build could not be checked against one."""
        sheet = xlsx.SheetOut(name="S", rows=[["a", 1], ["b", 2]])
        self.assertEqual(xlsx.write_workbook([sheet]),
                         xlsx.write_workbook([sheet]))

    def test_a_workbook_with_no_sheets_is_refused(self):
        with self.assertRaises(ValueError):
            xlsx.write_workbook([])

    def test_a_value_no_spreadsheet_can_hold_is_refused(self):
        """Writing a NaN produces a file that opens with a corrupt-content
        warning, which is worse than refusing to write it."""
        for bad in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(value=bad):
                with self.assertRaises(ValueError):
                    xlsx.write_workbook(
                        [xlsx.SheetOut(name="S", rows=[[bad]])])

    def test_a_formula_is_written_beside_a_value_not_instead_of_one(self):
        """The file opens with a correct number showing even in a reader that
        does not evaluate formulas, and a person can still see how it was
        derived."""
        sheet = xlsx.SheetOut(name="S", rows=[[2, 3, 6]],
                              formulas={(1, 3): "A1*B1"})
        book = xlsx.read_workbook(xlsx.write_workbook([sheet]))
        cell = book.first.rows[0][2]
        self.assertEqual(cell.value, 6)
        self.assertEqual(cell.formula, "A1*B1")

    def test_text_that_xml_cannot_carry_is_stripped_not_escaped_wrongly(self):
        sheet = xlsx.SheetOut(name="S", rows=[["a\x00b", "<&>\"'"]])
        book = xlsx.read_workbook(xlsx.write_workbook([sheet]))
        self.assertEqual(book.first.values()[0], ["ab", "<&>\"'"])

    def test_a_sheet_name_with_markup_in_it_does_not_break_the_file(self):
        sheet = xlsx.SheetOut(name='a&b<c>"d"', rows=[[1]])
        book = xlsx.read_workbook(xlsx.write_workbook([sheet]))
        self.assertEqual(book.first.name, 'a&b<c>"d"')


@NEEDS_OPENPYXL
class AnIndependentLibraryAgrees(unittest.TestCase):
    """Our own round-trip proves we agree with ourselves, which is worth very
    little. These ask another implementation."""

    def our_file(self):
        return xlsx.write_workbook([xlsx.SheetOut(
            name="Data",
            rows=[["item", "qty", "price", "total", "when"],
                  ["widget", 3, 2.5, 7.5, dt.date(2026, 3, 15)],
                  ["gadget", 10, 1.25, 12.5, dt.date(2026, 4, 1)],
                  ["TOTAL", None, None, 20.0, None]],
            formulas={(2, 4): "B2*C2", (3, 4): "B3*C3",
                      (4, 4): "SUM(D2:D3)"},
            number_formats={3: "0.00", 4: "0.00", 5: "yyyy-mm-dd"},
            freeze_below_row=2, widths={1: 18})])

    def test_openpyxl_loads_our_file_without_complaint(self):
        import warnings

        with warnings.catch_warnings():
            warnings.simplefilter("error")
            book = openpyxl.load_workbook(io.BytesIO(self.our_file()))
        self.assertEqual(book.sheetnames, ["Data"])

    def test_openpyxl_reads_the_values_we_meant(self):
        sheet = openpyxl.load_workbook(
            io.BytesIO(self.our_file()), data_only=False)["Data"]
        self.assertEqual([c.value for c in sheet[1]],
                         ["item", "qty", "price", "total", "when"])
        self.assertEqual(sheet["B2"].value, 3)
        self.assertEqual(sheet["C2"].value, 2.5)
        self.assertEqual(sheet["E2"].value, dt.datetime(2026, 3, 15))

    def test_openpyxl_sees_the_formulas_we_wrote(self):
        sheet = openpyxl.load_workbook(
            io.BytesIO(self.our_file()), data_only=False)["Data"]
        self.assertEqual(sheet["D2"].value, "=B2*C2")
        self.assertEqual(sheet["D4"].value, "=SUM(D2:D3)")

    def test_openpyxl_sees_the_number_formats_and_the_frozen_header(self):
        sheet = openpyxl.load_workbook(io.BytesIO(self.our_file()))["Data"]
        self.assertEqual(sheet["C2"].number_format, "0.00")
        self.assertEqual(sheet["E2"].number_format, "yyyy-mm-dd")
        self.assertEqual(sheet.freeze_panes, "A2")

    def test_we_read_what_openpyxl_writes(self):
        out = openpyxl.Workbook()
        ws = out.active
        ws.title = "Theirs"
        ws.append(["name", "amount", "when", "flag"])
        ws.append(["alpha", 12.75, dt.date(2024, 7, 4), True])
        ws.append(["beta", 3, dt.datetime(2024, 7, 5, 13, 30), False])
        ws["C2"].number_format = "yyyy-mm-dd"
        ws["C3"].number_format = "yyyy-mm-dd hh:mm"
        second = out.create_sheet("Second")
        second.append(["x"])
        second.append([99])
        buffer = io.BytesIO()
        out.save(buffer)

        book = xlsx.read_workbook(buffer.getvalue())
        self.assertEqual([s.name for s in book.sheets], ["Theirs", "Second"])
        rows = book.sheet("Theirs").values()
        self.assertEqual(rows[0][:4], ["name", "amount", "when", "flag"])
        self.assertEqual(rows[1][:4],
                         ["alpha", 12.75, dt.date(2024, 7, 4), True])
        self.assertEqual(rows[2][2], dt.datetime(2024, 7, 5, 13, 30))
        self.assertIs(rows[2][3], False)
        self.assertEqual(book.sheet("Second").values(), [["x"], [99]])

    def test_we_read_a_shared_string_table_written_by_another_library(self):
        """openpyxl uses the shared-string table where we use inline strings,
        so this exercises the read path our own writer never produces."""
        out = openpyxl.Workbook()
        ws = out.active
        for text in ("repeated", "repeated", "distinct"):
            ws.append([text])
        buffer = io.BytesIO()
        out.save(buffer)
        book = xlsx.read_workbook(buffer.getvalue())
        self.assertEqual(book.first.values(),
                         [["repeated"], ["repeated"], ["distinct"]])

    def test_a_formula_with_no_cached_value_reads_as_a_formula(self):
        """A library that writes formulas without evaluating them leaves no
        cached value. The formula must still be visible."""
        out = openpyxl.Workbook()
        ws = out.active
        ws["A1"] = 4
        ws["B1"] = "=A1*3"
        buffer = io.BytesIO()
        out.save(buffer)
        cell = xlsx.read_workbook(buffer.getvalue()).first.rows[0][1]
        self.assertEqual(cell.formula, "A1*3")
        self.assertTrue(cell.is_formula)


if __name__ == "__main__":
    unittest.main()
